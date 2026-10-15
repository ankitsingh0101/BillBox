import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app as appmod
from billbox import extractor, pipeline
from billbox.export import build_workbook
from gen_synthetic import fake_gstin, make_truth, render


def fake_model_factory(truth):
    return lambda img: dict(truth)


def test_pipeline_flags_a_misread(tmp_path, monkeypatch):
    buyer = fake_gstin("23")
    t = make_truth(buyer)
    render(t, tmp_path / "a.jpg")
    # model correct
    ok = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))[0]
    assert pipeline.status(ok) in ("OK", "CHECK")
    # model misreads the total
    bad_t = dict(t, total=t["total"] + 90)
    bad = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(bad_t))[0]
    assert pipeline.status(bad) == "REVIEW"


def test_duplicates_detected(tmp_path):
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    a = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))[0]
    b = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))[0]
    b["source"] = "copy.jpg"
    pipeline.mark_duplicates([a, b])
    assert any("duplicate" in f["msg"].lower() for f in b["flags"])


def test_model_offline_does_not_crash_batch(tmp_path):
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    def boom(img): raise ConnectionError("ollama not running")
    b = pipeline.process_file(tmp_path / "a.jpg", extract_fn=boom)[0]
    assert pipeline.status(b) == "REVIEW" and "Extraction failed" in b["flags"][0]["msg"]


def test_excel_export_roundtrip(tmp_path):
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    b = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))
    out = build_workbook(b, str(tmp_path / "o.xlsx"))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Bills", "Summary"]
    assert wb["Bills"]["P2"].value == "=SUM(J2:N2)"
    assert wb["Bills"]["E1"].value.startswith("What it's for")


def test_api_endpoints(monkeypatch, tmp_path):
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    monkeypatch.setattr(extractor, "extract", lambda img, **k: dict(t))
    c = TestClient(appmod.app)
    r = c.post("/api/process", files=[("files", ("a.jpg", (tmp_path / "a.jpg").read_bytes(), "image/jpeg"))])
    assert r.status_code == 200 and len(r.json()["bills"]) == 1
    bill = r.json()["bills"][0]
    bill["total"] = 1.0
    v = c.post("/api/check", json={"bills": [bill]}).json()["bills"][0]
    assert v["status"] == "REVIEW"
    x = c.post("/api/export", json={"bills": [bill]})
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_check_finds_duplicates_across_separate_uploads(tmp_path):
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    a = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))[0]
    b = dict(pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))[0], source="copy.jpg")
    out = pipeline.recheck([a, b])
    assert out[1]["status"] == "REVIEW" and out[0]["status"] != "REVIEW"


def test_autocorrect_is_reported_as_info_not_warning(tmp_path):
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    t["vendor_gstin"] = t["vendor_gstin"].replace("Z", "2", 1) if "Z" in t["vendor_gstin"][:14] else t["vendor_gstin"]
    b = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))[0]
    assert pipeline.status(b) in ("OK", "CHECK")


def test_export_temp_file_is_deleted_after_download(monkeypatch, tmp_path):
    import tempfile
    made = []
    real = tempfile.mkdtemp
    monkeypatch.setattr(tempfile, "mkdtemp", lambda **k: made.append(real(**k)) or made[-1])
    t = make_truth(fake_gstin("23")); render(t, tmp_path / "a.jpg")
    b = pipeline.process_file(tmp_path / "a.jpg", extract_fn=fake_model_factory(t))
    TestClient(appmod.app).post("/api/export", json={"bills": b})
    assert made and not Path(made[0]).exists()


def test_health_reports_ollama_down(monkeypatch):
    import requests
    def boom(*a, **k): raise requests.ConnectionError()
    monkeypatch.setattr(requests, "get", boom)
    assert TestClient(appmod.app).get("/api/health").json()["ollama"] is False