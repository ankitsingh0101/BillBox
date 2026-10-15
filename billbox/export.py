from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import status

FONT = "Arial"
KEYS = ["src", "st", "vendor", "cat", "sum", "vg", "bg", "inv", "date", "tax", "cgst", "sgst", "igst",
        "cess", "tot", "calc", "diff", "notes"]
HEAD = ["Source", "Status", "Vendor", "Category", "What it's for (AI summary)", "Vendor GSTIN", "Buyer GSTIN",
        "Invoice No", "Date", "Taxable (Rs)", "CGST (Rs)", "SGST (Rs)", "IGST (Rs)", "Cess (Rs)", "Total (Rs)",
        "Recalculated Total", "Difference", "Review notes"]
WIDTHS = [22, 9, 26, 20, 44, 18, 18, 16, 13, 13, 11, 11, 11, 10, 13, 15, 12, 60]
C = {k: get_column_letter(i) for i, k in enumerate(KEYS, start=1)}
FILL = {"REVIEW": "F8CBAD", "CHECK": "FFE699", "OK": None}
MONEY = "#,##0.00;(#,##0.00);-"


def build_workbook(bills: list[dict], path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, b in enumerate(bills, start=2):
        st = status(b)
        try:
            d = date.fromisoformat(b["invoice_date"])
        except (ValueError, TypeError):
            d = b.get("invoice_date") or ""
        notes = "; ".join(f["msg"] for f in b.get("flags", []))
        ws.append([b.get("source"), st, b.get("vendor_name"), b.get("category"), b.get("summary"),
                   b.get("vendor_gstin"), b.get("buyer_gstin"), b.get("invoice_number"), d,
                   b.get("taxable_value"), b.get("cgst"), b.get("sgst"), b.get("igst"), b.get("cess"),
                   b.get("total"),
                   f"=SUM({C['tax']}{i}:{C['cess']}{i})", f"={C['tot']}{i}-{C['calc']}{i}", notes])
        fill = FILL[st]
        for c in ws[i]:
            c.font = Font(name=FONT)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
        ws[f"{C['date']}{i}"].number_format = "dd-mmm-yyyy"
        for k in ("tax", "cgst", "sgst", "igst", "cess", "tot", "calc", "diff"):
            ws[f"{C[k]}{i}"].number_format = MONEY
        for k in ("sum", "notes"):
            ws[f"{C[k]}{i}"].alignment = Alignment(wrap_text=True, vertical="top")

    for idx, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    last = max(len(bills) + 1, 2)
    ws.auto_filter.ref = f"A1:{C['notes']}{last}"

    def rng(k):
        return f"Bills!{C[k]}2:{C[k]}{last}"

    s = wb.create_sheet("Summary")
    rows = [
        ("BillBox summary", None),
        ("Bills processed", f"=COUNTA({rng('src')})"),
        ("OK (passed every check)", f'=COUNTIF({rng("st")},"OK")'),
        ("CHECK (warnings)", f'=COUNTIF({rng("st")},"CHECK")'),
        ("REVIEW (errors - fix before filing)", f'=COUNTIF({rng("st")},"REVIEW")'),
        (None, None),
        ("Total taxable value (Rs)", f"=SUM({rng('tax')})"),
        ("Total CGST (Rs)", f"=SUM({rng('cgst')})"),
        ("Total SGST (Rs)", f"=SUM({rng('sgst')})"),
        ("Total IGST (Rs)", f"=SUM({rng('igst')})"),
        ("Total invoice value (Rs)", f"=SUM({rng('tot')})"),
        ("Input tax in OK bills only (Rs)",
         "=" + "+".join(f'SUMIFS({rng(k)},{rng("st")},"OK")' for k in ("cgst", "sgst", "igst"))),
        (None, None),
        ("Note", "Status colours: red = needs a human, amber = look once, none = passed all checks. "
                 "Checks: GSTIN checksum, tax maths (+/- Rs 1 round-off), CGST=SGST, standard GST slabs, duplicates. "
                 "Category and 'What it's for' are written by the AI model from the bill text and are not verified by code."),
    ]
    for r in rows:
        s.append(list(r))
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            if c.column == 2 and isinstance(c.value, str) and c.value.startswith("="):
                c.number_format = MONEY if c.row >= 7 else "0"
    s["B14"].alignment = Alignment(wrap_text=True, vertical="top")
    s.column_dimensions["A"].width = 38
    s.column_dimensions["B"].width = 70
    wb.save(path)
    return path